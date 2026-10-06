import os
import re
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.worksheet.filters import FilterColumn
from Services.Logger import Logger

logger = Logger.instance()

# 课次状态 -> 输出字母:看课(直/录) 记 T,到记 D,缺记 F,销记 X
COURSE_LETTERS = {"直": "T", "录": "T", "缺": "F", "到": "D", "销": "X"}
# 作业状态 -> 输出字母:交了作业(作) 记 T,其余记 F
HOMEWORK_LETTERS = {"作": "T"}
# 完课只认这一个字母,「计数」与全勤作业表的筛选都按它算
DONE_LETTER = "T"


class ExcelExportService:
    """负责将 data.xlsx 解析拆分为课程看课表、作业提交表、全勤作业表、上次活跃时间表并设置样式导出的服务类"""

    @staticmethod
    def parse_course(val: str) -> str:
        """解析课程部分：直/录 -> T，到 -> D，缺 -> F，销 -> X（空值按 F 处理）"""
        if pd.isna(val):
            return "F"
        parts = str(val).split("/")
        return COURSE_LETTERS.get(parts[0], "F")

    @staticmethod
    def parse_homework(val: str) -> str:
        """解析作业部分：提交作业(作) -> T，其余 -> F"""
        if pd.isna(val):
            return "F"
        parts = str(val).split("/")
        return HOMEWORK_LETTERS.get(parts[1], "F") if len(parts) > 1 else "F"

    def process_df(self, df: pd.DataFrame, target_cols: list, parse_func, count_col_name: str) -> pd.DataFrame:
        """通用处理函数：状态映射 -> 统计完课次数（只数 T） -> 插入计数列"""
        target = df[target_cols]
        mapper = target.map if hasattr(target, "map") else target.applymap

        letter_df = mapper(parse_func)
        counts = letter_df.eq(DONE_LETTER).sum(axis=1)

        df[target_cols] = letter_df

        name_idx = df.columns.get_loc("真实姓名")
        df.insert(loc=name_idx + 1, column=count_col_name, value=counts)

        return df

    @staticmethod
    def highlight_failed(val):
        """不是完课(T)的单元格背景标红:F 缺席、D 到课、X 销课"""
        if val != DONE_LETTER:
            return "background-color: #FFCCCC; color: red;"
        return ""

    def save_styled_excel(self, styled_data, file_path: str):
        """将 Styler 对象写入 Excel 并统一进行 openpyxl 样式设置"""
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        center_alignment = Alignment(horizontal="center", vertical="center")

        with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
            styled_data.to_excel(writer, index=False)
            ws = writer.sheets["Sheet1"]

            target_filter_names = {"学号", "计数", "真实姓名"}
            headers = {}

            for col in ws.iter_cols():
                col_letter = col[0].column_letter
                original_header_val = str(col[0].value or "").strip()
                header_val = original_header_val

                match = re.search(r"第(\d+)节课", header_val)
                if match:
                    header_val = match.group(1)
                    col[0].value = header_val

                col_idx = col[0].column
                headers[original_header_val] = col_idx
                headers[header_val] = col_idx

                for cell in col:
                    cell.alignment = center_alignment

                if header_val.isdigit():
                    ws.column_dimensions[col_letter].width = 4.5
                else:
                    max_len = 0
                    for cell in col:
                        if cell.value is not None:
                            s = str(cell.value)
                            cell_len = sum(2 if ord(c) > 127 else 1 for c in s)
                            if cell_len > max_len:
                                max_len = cell_len
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

            target_indices = [headers[col] for col in target_filter_names if col in headers]

            if target_indices:
                max_row = ws.max_row
                min_col_idx = min(target_indices)
                max_col_idx = max(target_indices)

                ws.auto_filter.ref = (
                    f"{ws.cell(row=1, column=min_col_idx).coordinate}:"
                    f"{ws.cell(row=max_row, column=max_col_idx).coordinate}"
                )

                ws.auto_filter.filterColumn.clear()

                for col_idx in range(min_col_idx, max_col_idx + 1):
                    col_id = col_idx - min_col_idx
                    is_hidden = col_idx not in target_indices
                    col_filter = FilterColumn(colId=col_id, hiddenButton=is_hidden)
                    ws.auto_filter.filterColumn.append(col_filter)

    def export_split_tables(self, record_path: str) -> tuple[bool, str]:
        """执行拆分并保存导出的 Excel 文件至当前记录路径"""
        data_excel_path = os.path.join(record_path, "data.xlsx")
        if not os.path.exists(data_excel_path):
            logger.error(f"未找到源文件: {data_excel_path}")
            return False, f"未找到源文件: {data_excel_path}"

        try:
            data = pd.read_excel(data_excel_path)
            data.columns = data.columns.str.strip()

            if "学员学籍状态" in data.columns:
                data = data[data["学员学籍状态"] == "在读"]

            expected_lessons = [f"第{i}节课" for i in range(1, 33)]
            lesson_cols = [col for col in expected_lessons if col in data.columns]
            select_cols = ["学号", "真实姓名"] + lesson_cols

            data = data[select_cols].dropna(axis=1, how="all")

            target_cols = [col for col in data.columns if col not in ("学号", "真实姓名")]
            max_lessons = len(target_cols)

            course_df = self.process_df(
                data.copy(), target_cols, self.parse_course, count_col_name="计数"
            )
            homework_df = self.process_df(
                data.copy(), target_cols, self.parse_homework, count_col_name="计数"
            )

            completed_student_ids = course_df[course_df["计数"] == max_lessons]["学号"]
            finished_course_homework = homework_df[
                homework_df["学号"].isin(completed_student_ids)
            ].copy()

            course_styled = course_df.style.map(self.highlight_failed, subset=target_cols)
            homework_styled = homework_df.style.map(self.highlight_failed, subset=target_cols)
            finished_course_homework_styled = finished_course_homework.style.map(
                self.highlight_failed, subset=target_cols
            )

            split_dir = os.path.join(record_path, "split")
            course_path = os.path.join(split_dir, "course.xlsx")
            homework_path = os.path.join(split_dir, "homework.xlsx")
            finished_path = os.path.join(split_dir, "finished_course_homework.xlsx")

            self.save_styled_excel(course_styled, course_path)
            self.save_styled_excel(homework_styled, homework_path)
            self.save_styled_excel(finished_course_homework_styled, finished_path)

            # 上次活跃时间表(3 字段:学号 / 真实姓名 / 上次活跃时间)。
            # 它依赖班级全部历史记录,算不出来时只记错误,不影响其余三张表。
            try:
                from Services.LastActiveService import LastActiveService

                active_map = LastActiveService().compute(record_path)
                active_rows = []
                for _, row in data.iterrows():
                    sid = "" if pd.isna(row["学号"]) else str(row["学号"]).strip()
                    if not sid:
                        continue
                    name = (
                        ""
                        if pd.isna(row["真实姓名"])
                        else str(row["真实姓名"]).strip()
                    )
                    active_rows.append(
                        {
                            "学号": sid,
                            "真实姓名": name,
                            "上次活跃时间": active_map.get(sid, ""),
                        }
                    )
                active_df = pd.DataFrame(
                    active_rows, columns=["学号", "真实姓名", "上次活跃时间"]
                )
                self.save_styled_excel(
                    active_df.style, os.path.join(split_dir, "last_active.xlsx")
                )
            except Exception as e:
                logger.error(f"上次活跃时间表生成失败(不影响其余拆分表): {e}")

            return True, "拆分导出表格成功！"
        except PermissionError:
            logger.error("导出失败:目标 Excel 文件被占用,请关闭后重试!")
            return False, "导出失败:目标 Excel 文件被占用,请关闭后重试!"
        except Exception as e:
            logger.error(f"拆分导出时发生错误: {e}")
            return False, f"拆分导出时发生错误: {e}"