import os
import re
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.worksheet.filters import FilterColumn
from Services.Logger import Logger

logger = Logger.instance()


def _parse_course(val: str) -> bool:
    """解析课程部分:是否包含 '直' 或 '录'"""
    if pd.isna(val):
        return False
    parts = str(val).split("/")
    return parts[0] in ("直", "录") if len(parts) > 0 else False


def _parse_homework(val: str) -> bool:
    """解析作业部分:是否包含 '作'"""
    if pd.isna(val):
        return False
    parts = str(val).split("/")
    return parts[1] in ("作") if len(parts) > 1 else False


def _process_df(
    df: pd.DataFrame, target_cols: list, parse_func, count_col_name: str
) -> pd.DataFrame:
    """通用处理函数:布尔映射 -> 统计完成次数 -> 插入计数列 -> 替换 False 为F"""
    target = df[target_cols]
    mapper = target.map if hasattr(target, "map") else target.applymap

    bool_df = mapper(parse_func)

    counts = bool_df.sum(axis=1)

    formatted_df = bool_df.astype(object)
    df[target_cols] = formatted_df.where(bool_df, "F")
    df[target_cols] = df[target_cols].where(~bool_df, "T")

    name_idx = df.columns.get_loc("真实姓名")
    df.insert(loc=name_idx + 1, column=count_col_name, value=counts)

    return df


def _highlight_empty(val):
    """F单元格背景标红"""
    if val == "F" or pd.isna(val):
        return "background-color: #FFCCCC; color: red;"
    return ""


def _save_styled_excel(styled_data, file_path: str):
    """将 Styler 对象写入 Excel 并统一进行 openpyxl 样式设置"""

    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    center_alignment = Alignment(horizontal="center", vertical="center")

    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        styled_data.to_excel(writer, index=False)
        ws = writer.sheets["Sheet1"]

        target_filter_names = {"学号", "计数", "真实姓名"}
        headers = {}

        for col in ws.iter_cols():
            col_letter = col[0].column_letter
            original_header_val = str(col[0].value or "").strip()
            header_val = original_header_val

            match = re.search(r"第(\d+)节课", header_val)
            if match:
                header_val = match.group(1)
                col[0].value = header_val

            col_idx = col[0].column
            headers[original_header_val] = col_idx
            headers[header_val] = col_idx

            for cell in col:
                cell.alignment = center_alignment

            if header_val.isdigit():
                ws.column_dimensions[col_letter].width = 4.5
            else:
                max_len = 0
                for cell in col:
                    if cell.value is not None:
                        s = str(cell.value)
                        cell_len = sum(2 if ord(c) > 127 else 1 for c in s)
                        if cell_len > max_len:
                            max_len = cell_len
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        target_indices = [headers[col] for col in target_filter_names if col in headers]

        if target_indices:
            max_row = ws.max_row
            min_col_idx = min(target_indices)
            max_col_idx = max(target_indices)

            ws.auto_filter.ref = (
                f"{ws.cell(row=1, column=min_col_idx).coordinate}:"
                f"{ws.cell(row=max_row, column=max_col_idx).coordinate}"
            )

            ws.auto_filter.filterColumn.clear()

            for col_idx in range(min_col_idx, max_col_idx + 1):
                col_id = col_idx - min_col_idx
                is_hidden = col_idx not in target_indices
                col_filter = FilterColumn(colId=col_id, hiddenButton=is_hidden)
                ws.auto_filter.filterColumn.append(col_filter)


class ExcelSplitterService:
    def split_data(self, record_dir_path: str) -> bool:
        data_excel_path = os.path.join(record_dir_path, "data.xlsx")
        if not os.path.exists(data_excel_path):
            logger.error(f"未找到 {data_excel_path}")
            return False

        try:
            data = pd.read_excel(data_excel_path)
        except PermissionError:
            logger.error(f"读写 Excel 权限被拒绝,请检查文件是否被占用: {data_excel_path}")
            return False
        except Exception as e:
            logger.error(f"读取 Excel 失败: {e}")
            return False

        data.columns = data.columns.str.strip()

        if "学员学籍状态" not in data.columns:
            logger.error("找不到 '学员学籍状态' 列")
            return False
        data = data[data["学员学籍状态"] == "在读"]

        expected_lessons = [f"第{i}节课" for i in range(1, 33)]
        lesson_cols = [col for col in expected_lessons if col in data.columns]

        if "学号" not in data.columns or "真实姓名" not in data.columns:
            logger.error("缺少 '学号' 或 '真实姓名' 列")
            return False

        select_cols = ["学号", "真实姓名"] + lesson_cols
        data = data[select_cols].dropna(axis=1, how="all")

        target_cols = [col for col in data.columns if col not in ("学号", "真实姓名")]
        max_lessons = len(target_cols)

        course_df = _process_df(
            data.copy(),
            target_cols,
            _parse_course,
            count_col_name="计数",
        )
        homework_df = _process_df(
            data.copy(),
            target_cols,
            _parse_homework,
            count_col_name="计数",
        )

        completed_student_ids = course_df[course_df["计数"] == max_lessons]["学号"]
        finished_course_homework = homework_df[
            homework_df["学号"].isin(completed_student_ids)
        ].copy()

        course_styled = course_df.style.map(_highlight_empty, subset=target_cols)
        homework_styled = homework_df.style.map(_highlight_empty, subset=target_cols)
        finished_course_homework_styled = finished_course_homework.style.map(
            _highlight_empty, subset=target_cols
        )

        course_file = os.path.join(record_dir_path, "split", "course.xlsx")
        finished_file = os.path.join(record_dir_path, "split", "finished_course_homework.xlsx")
        homework_file = os.path.join(record_dir_path, "split", "homework.xlsx")

        try:
            _save_styled_excel(course_styled, course_file)
            logger.info(f"已保存 {course_file}")
            _save_styled_excel(homework_styled, homework_file)
            logger.info(f"已保存 {homework_file}")
            _save_styled_excel(finished_course_homework_styled, finished_file)
            logger.info(f"已保存 {finished_file}")
        except PermissionError:
            logger.error("保存 Excel 时权限被拒绝,请检查文件是否被占用")
            return False
        except Exception as e:
            logger.error(f"保存 Excel 失败: {e}")
            return False

        return True